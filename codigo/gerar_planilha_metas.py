# Gera relatorios/plano_metas_lojas.xlsx a partir de dados/processados/insumo_plano_metas.csv (notebook 06).
# Os valores calculados ficam em fórmulas; abra no Excel ou recalcule antes de ler com pandas.
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.comments import Comment

RAIZ = Path(__file__).resolve().parents[1]
insumo = pd.read_csv(RAIZ / 'dados' / 'processados' / 'insumo_plano_metas.csv')
n = len(insumo); ult = n + 1
FONTE = 'Arial'
f_normal = Font(name=FONTE, size=10)
f_input = Font(name=FONTE, size=10, color='0000FF')
f_link = Font(name=FONTE, size=10, color='008000')
f_cab = Font(name=FONTE, size=10, bold=True, color='FFFFFF')
f_tit = Font(name=FONTE, size=13, bold=True)
f_neg = Font(name=FONTE, size=10, bold=True)
fill_cab = PatternFill('solid', start_color='1F4E78')
fill_chave = PatternFill('solid', start_color='FFFF00')
borda = Border(bottom=Side(style='thin', color='BFBFBF'))
PCT, RS, NUM = '0.0%;(0.0%);-', '"R$" #,##0;("R$" #,##0);-', '#,##0.0'

wb = Workbook()
par = wb.active; par.title = 'Parâmetros'
met = wb.create_sheet('Metas por Loja')
res = wb.create_sheet('Resumo por Cluster')
pre = wb.create_sheet('Premissas')

par['A1'] = 'Plano de metas por loja — parâmetros'; par['A1'].font = f_tit
par['A2'] = 'Células em azul são entradas editáveis. As demais são calculadas.'; par['A2'].font = Font(name=FONTE, size=9, italic=True)
entradas = [('Meta de manutenção (loja na mediana do cluster ou acima)', 0.02, 'Regra definida no notebook 06: quem já é referência do grupo recebe só manutenção.'),
            ('Captura do gap até a mediana do cluster', 0.25, 'Parcela do gap somada à meta das lojas abaixo da mediana.'),
            ('Gap mínimo para prioridade "Ação"', 0.10, 'Abaixo desse gap a loja fica em "Acompanhamento".'),
            ('Compromisso de crescimento da rede', 0.05, 'Compromisso assinado pela diretoria com o conselho (+5%).'),
            ('Teto de meta por loja', 0.20, 'Nenhuma loja recebe meta acima deste valor; o fator de calibração compensa o que o teto corta.')]
for i, (rot, val, nota) in enumerate(entradas, start=4):
    par[f'A{i}'] = rot; par[f'A{i}'].font = f_normal
    c = par[f'B{i}']; c.value = val; c.font = f_input; c.number_format = PCT; c.comment = Comment(nota, 'Plano de metas')
par['B7'].fill = fill_chave
par['B8'].fill = fill_chave
J = "'Metas por Loja'!$J$2:$J$" + str(ult); K = "'Metas por Loja'!$K$2:$K$" + str(ult); G = "'Metas por Loja'!$G$2:$G$" + str(ult)
calc = [('Crescimento ponderado com as metas iniciais', f"=SUMPRODUCT({J},{K})/SUM({J})", PCT),
        ('Fator sem teto (iteração 0)', '=B7/B11', '0.0000')]
for it in range(1, 7):
    ant = f'B{11 + it}'
    calc.append((f'Fator com teto (iteração {it})',
                 f"=($B$7*SUM({J})-$B$8*SUMPRODUCT({J},--({K}*{ant}>$B$8)))/SUMPRODUCT({J},{K},--({K}*{ant}<=$B$8))", '0.0000'))
calc += [('Fator de calibração final', '=B18', '0.0000'),
         ('Convergiu? (iterações 5 e 6 iguais)', '=IF(ABS(B18-B17)<0.0000001,"Sim","Não")', '@'),
         ('Meta de manutenção após calibração', '=B4*B19', PCT),
         ('Lojas no teto', f"=SUMPRODUCT(--({K}*B19>B8))", '0'),
         ('Média simples das metas iniciais', f"=AVERAGE({K})", PCT),
         ('Média simples das metas finais', f"=AVERAGE({G})", PCT),
         ('Crescimento ponderado com as metas finais (checagem)', f"=SUMPRODUCT({J},{G})/SUM({J})", PCT),
         ('Faturamento anual estimado da rede (R$)', f"=SUM({J})", RS),
         ('Ganho anual previsto com as metas finais (R$)', '=B25*B26', RS)]
for i, (rot, form, fmt) in enumerate(calc, start=11):
    par[f'A{i}'] = rot; par[f'A{i}'].font = f_normal
    c = par[f'B{i}']; c.value = form; c.font = f_normal; c.number_format = fmt
par['A10'] = 'Calibração'; par['A10'].font = f_neg

par['A31'] = 'Benchmark por cluster (mediana da venda média por dia)'; par['A18'].font = f_neg
for j, cab in enumerate(['Cluster', 'Nome do Cluster', 'Benchmark do Grupo (R$/dia)'], start=1):
    c = par.cell(row=32, column=j, value=cab); c.font = f_cab; c.fill = fill_cab
nomes = insumo.drop_duplicates('cluster').sort_values('cluster')[['cluster', 'nome_cluster']].values
for k, (cl, nome) in enumerate(nomes):
    r = 33 + k
    par[f'A{r}'] = int(cl); par[f'A{r}'].font = f_normal
    par[f'B{r}'] = nome; par[f'B{r}'].font = f_normal
    ref = f'C{r}'
    par[ref] = ArrayFormula(ref, f"=MEDIAN(IF('Metas por Loja'!$B$2:$B${ult}=A{r},'Metas por Loja'!$D$2:$D${ult}))")
    par[ref].font = f_normal; par[ref].number_format = RS
par.column_dimensions['A'].width = 58; par.column_dimensions['B'].width = 36; par.column_dimensions['C'].width = 26

cabecalhos = ['Loja', 'Cluster', 'Nome do Cluster', 'Venda Média por Dia (R$)', 'Benchmark do Grupo (R$/dia)', 'Gap Percentual',
              'Meta Percentual', 'Prioridade', 'Dias de Operação por Ano', 'Faturamento Anual (R$)', 'Meta Inicial (antes da calibração)',
              'Meta de Venda por Dia (R$)', 'Abre aos Domingos', 'Permanência no Cluster', 'Observação']
for j, cab in enumerate(cabecalhos, start=1):
    c = met.cell(row=1, column=j, value=cab); c.font = f_cab; c.fill = fill_cab
    c.alignment = Alignment(wrap_text=True, vertical='center')
met['F1'].comment = Comment('Crescimento necessário para a loja chegar à mediana do cluster: (benchmark - venda) / venda. Zero para quem está na mediana ou acima.', 'Plano de metas')
met['I1'].comment = Comment('Proporção de dias abertos da loja no histórico x 365 (notebook 02/03).', 'Plano de metas')
for i, linha in enumerate(insumo.itertuples(index=False), start=2):
    met[f'A{i}'] = int(linha.loja); met[f'B{i}'] = int(linha.cluster); met[f'C{i}'] = linha.nome_cluster
    met[f'D{i}'] = float(linha.venda_media_dia); met[f'I{i}'] = float(linha.dias_operacao_ano)
    met[f'E{i}'] = f'=INDEX(Parâmetros!$C$33:$C$37,MATCH(B{i},Parâmetros!$A$33:$A$37,0))'
    met[f'F{i}'] = f'=MAX(0,(E{i}-D{i})/D{i})'
    met[f'K{i}'] = f'=Parâmetros!$B$4+Parâmetros!$B$5*F{i}'
    met[f'G{i}'] = f'=MIN(K{i}*Parâmetros!$B$19,Parâmetros!$B$8)'
    met[f'H{i}'] = f'=IF(F{i}>=Parâmetros!$B$6,"Ação","Acompanhamento")'
    met[f'M{i}'] = int(linha.abre_domingo); met[f'N{i}'] = float(linha.permanencia_cluster)
    met[f'O{i}'] = f'=IF(AND(B{i}=4,M{i}=1),IF(N{i}<0.8,"Subgrupo 7 dias; fronteira de cluster","Subgrupo 7 dias"),IF(N{i}<0.8,"Fronteira de cluster",""))'
    met[f'J{i}'] = f'=D{i}*I{i}'
    met[f'L{i}'] = f'=D{i}*(1+G{i})'
    for col in 'ABCDEFGHIJKLMNO':
        met[f'{col}{i}'].font = f_input if col in 'ABCDIMN' else (f_link if col == 'E' else f_normal)
    for col, fmt in {'D': RS, 'E': RS, 'F': PCT, 'G': PCT, 'I': NUM, 'J': RS, 'K': PCT, 'L': RS, 'N': PCT}.items():
        met[f'{col}{i}'].number_format = fmt
larg = [7, 8, 34, 16, 17, 11, 11, 16, 13, 17, 16, 16, 11, 13, 34]
for j, w in enumerate(larg, start=1):
    met.column_dimensions[met.cell(row=1, column=j).column_letter].width = w
met.row_dimensions[1].height = 42
met.freeze_panes = 'D2'
met.auto_filter.ref = f'A1:O{ult}'
met['G1'].comment = Comment('MIN(meta inicial x fator de calibração com teto, teto). A rede soma exatamente o compromisso ponderado pelo faturamento.', 'Plano de metas')
met['N1'].comment = Comment('Proporção das 100 subamostras de 80% (teste de estabilidade, notebook 05) em que a loja ficou no mesmo cluster.', 'Plano de metas')
met['O1'].comment = Comment('Subgrupo 7 dias: cluster 4 e abre aos domingos. Fronteira de cluster: permanência abaixo de 80%. Revisar a meta dessas lojas com o gestor.', 'Plano de metas')

res['A1'] = 'Resumo do plano por cluster'; res['A1'].font = f_tit
cab_r = ['Cluster', 'Nome do Cluster', 'Lojas', 'Lojas em Ação', 'Benchmark (R$/dia)', 'Meta Média', 'Meta Máxima', 'Crescimento Ponderado', 'Faturamento Anual (R$)']
for j, cab in enumerate(cab_r, start=1):
    c = res.cell(row=3, column=j, value=cab); c.font = f_cab; c.fill = fill_cab; c.alignment = Alignment(wrap_text=True)
M = "'Metas por Loja'!"
for k in range(5):
    r = 4 + k
    res[f'A{r}'] = f'=Parâmetros!A{33 + k}'; res[f'B{r}'] = f'=Parâmetros!B{33 + k}'
    res[f'C{r}'] = f'=COUNTIF({M}$B$2:$B${ult},A{r})'
    res[f'D{r}'] = f'=COUNTIFS({M}$B$2:$B${ult},A{r},{M}$H$2:$H${ult},"Ação")'
    res[f'E{r}'] = f'=Parâmetros!C{33 + k}'
    res[f'F{r}'] = f'=AVERAGEIF({M}$B$2:$B${ult},A{r},{M}$G$2:$G${ult})'
    res[f'G{r}'] = ArrayFormula(f'G{r}', f'=MAX(IF({M}$B$2:$B${ult}=A{r},{M}$G$2:$G${ult}))')
    res[f'H{r}'] = f'=SUMPRODUCT(({M}$B$2:$B${ult}=A{r})*{M}$J$2:$J${ult}*{M}$G$2:$G${ult})/I{r}'
    res[f'I{r}'] = f'=SUMIF({M}$B$2:$B${ult},A{r},{M}$J$2:$J${ult})'
res['A9'] = 'Rede'; res['B9'] = '—'
res['C9'] = '=SUM(C4:C8)'; res['D9'] = '=SUM(D4:D8)'; res['E9'] = f'=MEDIAN({M}D2:D{ult})'
res['F9'] = f'=AVERAGE({M}G2:G{ult})'; res['G9'] = f'=MAX({M}G2:G{ult})'; res['H9'] = '=Parâmetros!B25'; res['I9'] = '=SUM(I4:I8)'
for r in range(4, 10):
    for col, fmt in {'E': RS, 'F': PCT, 'G': PCT, 'H': PCT, 'I': RS}.items():
        res[f'{col}{r}'].number_format = fmt
    for col in 'ABCDEFGHI':
        res[f'{col}{r}'].font = f_neg if r == 9 else f_normal
        res[f'{col}{r}'].border = borda

res['A12'] = 'Distribuição das metas finais'; res['A12'].font = f_neg
for j, cab in enumerate(['Faixa', 'De', 'Até', 'Lojas', '% das lojas'], start=1):
    c = res.cell(row=13, column=j, value=cab); c.font = f_cab; c.fill = fill_cab
faixas = [('Manutenção', 0, '=Parâmetros!B21+0.00001'), ('Manutenção a 5%', '=C14', 0.05), ('5% a 10%', 0.05, 0.10),
          ('10% a 15%', 0.10, 0.15), ('15% a 20%', 0.15, '=Parâmetros!B8-0.00001'), ('No teto', '=C18', 1)]
for k, (rot, de, ate) in enumerate(faixas):
    r = 14 + k
    res[f'A{r}'] = rot; res[f'B{r}'] = de; res[f'C{r}'] = ate
    res[f'D{r}'] = f'=COUNTIFS({M}$G$2:$G${ult},">"&B{r},{M}$G$2:$G${ult},"<="&C{r})'
    res[f'E{r}'] = f'=D{r}/COUNT({M}$G$2:$G${ult})'
    for col in 'ABCDE':
        res[f'{col}{r}'].font = f_normal
    res[f'B{r}'].number_format = PCT; res[f'C{r}'].number_format = PCT; res[f'E{r}'].number_format = PCT
res['A20'] = 'Total'; res['D20'] = '=SUM(D14:D19)'; res['E20'] = '=SUM(E14:E19)'; res['E20'].number_format = PCT
res['A20'].font = f_neg; res['D20'].font = f_neg; res['E20'].font = f_neg
for col, w in zip('ABCDEFGHI', [16, 36, 8, 10, 16, 11, 11, 13, 20]):
    res.column_dimensions[col].width = w

textos = ['Premissas do plano de metas',
          '',
          '1. Referência de desempenho: venda média por dia de operação (dias abertos com faturamento > 0), histórico de 01/2013 a 07/2015.',
          '2. Benchmark do grupo: mediana da venda média por dia das lojas do mesmo cluster (K-Means, k = 5).',
          '3. Gap percentual: crescimento que a loja precisa para chegar ao benchmark = (benchmark - venda) / venda; zero para quem está na mediana ou acima.',
          '4. Meta inicial: manutenção para quem está na mediana ou acima; manutenção + captura x gap para quem está abaixo.',
          '5. Calibração: todas as metas multiplicadas pelo mesmo fator, com teto por loja (padrão 20%). O fator é recalculado por iteração para compensar o que o teto corta, e a rede cresce exatamente o compromisso ponderado pelo faturamento anual.',
          '6. Faturamento anual da loja: venda média por dia x dias de operação por ano (proporção de dias abertos no histórico x 365).',
          '7. Prioridade: "Ação" para lojas com gap igual ou acima do limiar; "Acompanhamento" para as demais.',
          '8. As metas são de faturamento, não de margem; não consideram custo de promoção, canibalização ou mudança de preço.',
          '9. Observação: "Subgrupo 7 dias" marca as lojas do cluster 4 que abrem aos domingos (perfil de conveniência, ticket bem abaixo do resto do cluster); "Fronteira de cluster" marca as lojas que ficaram no próprio cluster em menos de 80% das subamostras do teste de estabilidade. A meta dessas lojas deve ser revisada com o gestor.',
          '10. Para mudar a regra, altere as células em azul da aba Parâmetros; toda a planilha é recalculada.']
for i, t in enumerate(textos, start=1):
    pre[f'A{i}'] = t; pre[f'A{i}'].font = f_tit if i == 1 else f_normal
pre.column_dimensions['A'].width = 150

saida = RAIZ / 'relatorios' / 'plano_metas_lojas.xlsx'
wb.save(saida); print('salvo', saida)